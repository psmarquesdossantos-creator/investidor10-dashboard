import pdfplumber
import openpyxl
import json
import re
from datetime import datetime
import os

# ===== PARSE PDF RENDA FIXA =====
renda_fixa_positions = []
with pdfplumber.open('extrato-renda-fixa.pdf') as pdf:
    for page in pdf.pages:
        text = page.extract_text()
        if not text:
            continue
        lines = text.split('\n')
        for line in lines:
            # Match the CDB lines
            match = re.search(r'(\d{7,})\s+(\d{2}/\d{2}/\d{4})\s+(\d{2}/\d{2}/\d{4}).*?R\$\s*([\d.,]+)\s+CDB\s+([\d%]+\s+do\s+CDI).*?R\$\s*([\d.,]+).*?R\$\s*([\d.,]+).*?R\$\s*([\d.,]+).*?R\$\s*([\d.,]+).*?R\$\s*([\d.,]+).*?R\$\s*([\d.,]+).*?(\d+)%', line)
            if match:
                nota, inicio, vencimento, valor_aplicado, taxa, rendimento, retirada, desconto, bruto, liquido, ir = match.groups()
                renda_fixa_positions.append({
                    'nota': nota,
                    'data_inicio': inicio,
                    'vencimento': vencimento,
                    'valor_aplicado': float(valor_aplicado.replace('.', '').replace(',', '.')),
                    'taxa': taxa,
                    'rendimento': float(rendimento.replace('.', '').replace(',', '.') if rendimento else 0),
                    'retirada': float(retirada.replace('.', '').replace(',', '.') if retirada else 0),
                    'desconto': float(desconto.replace('.', '').replace(',', '.') if desconto else 0),
                    'bruto': float(bruto.replace('.', '').replace(',', '.') if bruto else 0),
                    'liquido': float(liquido.replace('.', '').replace(',', '.') if liquido else 0),
                    'ir_pct': int(ir)
                })
            # Also parse the LIG LIQUIDEZ
            match2 = re.match(r'(\d{7})\s+(\d{4}-\d{2}-\d{2})\s+(\d{4}-\d{2}-\d{2})\s+R\$\s*([\d.,]+)\s+LIG\s+(\d+%\s+do\s+CDI)', line)
            if match2:
                nota, inicio, vencimento, valor, taxa = match2.groups()
                renda_fixa_positions.append({
                    'nota': nota,
                    'data_inicio': inicio,
                    'vencimento': vencimento,
                    'valor_aplicado': float(valor.replace('.', '').replace(',', '.')),
                    'taxa': taxa,
                    'tipo': 'LIG'
                })

print(f'Renda Fixa positions parsed: {len(renda_fixa_positions)}')

# ===== PARSE PDF RENDA VARIÁVEL =====
renda_var_positions = []
with pdfplumber.open('extrato-renda-variavel.pdf') as pdf:
    for page in pdf.pages:
        text = page.extract_text()
        if not text:
            continue
        lines = text.split('\n')
        i = 0
        while i < len(lines) - 1:
            line = lines[i].strip()
            if re.match(r'^[A-Z]{4}\d{1,2}$', line):
                ticker = line
                if i + 1 < len(lines):
                    data_line = lines[i + 1]
                    nums = re.findall(r'[\d.,]+', lines[i + 1])
                    if len(nums) >= 7:
                        renda_var_positions.append({
                            'ticker': ticker,
                            'quantidade': int(nums[0].replace('.', '')),
                            'preco_medio': float(nums[1].replace('.', '').replace(',', '.')),
                            'custo_medio': float(nums[2].replace('.', '').replace(',', '.')),
                            'preco_atual': float(nums[3].replace('.', '').replace(',', '.')),
                            'valor_mercado': float(nums[4].replace('.', '').replace(',', '.')),
                            'retorno_pct': float(nums[5].replace(',', '.')),
                            'retorno_valor': float(nums[6].replace('.', '').replace(',', '.'))
                        })
                    i += 2
                    continue
            i += 1

print(f'Renda Variável positions parsed: {len(renda_var_positions)}')

# ===== PARSE XLSX MOVIMENTAÇÃO =====
movimentacoes = []
for fname in ['movimentacao-2026-10-10-18-50-38.xlsx', 'movimentacao-2026-10-10-18-51-16.xlsx']:
    wb = openpyxl.load_workbook(fname, data_only=True)
    ws = wb['Movimentação']
    for row in ws.iter_rows(min_row=2, values_only=True):
        entrada_saida, data, mov, produto, instituicao, qtd, preco, valor = row
        if data and mov:
            try:
                data_obj = datetime.strptime(str(data), '%d/%m/%Y')
            except:
                continue
            # Handle "-" in preco
            preco_val = 0
            if preco and str(preco) != '-':
                preco_val = float(preco)
            qtd_val = float(qtd) if qtd else 0
            valor_val = float(valor) if valor else 0
            movimentacoes.append({
                'entrada_saida': entrada_saida,
                'data': data_obj,
                'movimentacao': mov,
                'produto': produto,
                'instituicao': instituicao,
                'quantidade': qtd_val,
                'preco': preco_val,
                'valor': valor_val
            })

print(f'Movimentações parsed: {len(movimentacoes)}')

# ===== PARSE XLSX POSIÇÃO =====
posicoes_acoes = []
wb = openpyxl.load_workbook('posicao-2026-10-10-18-49-30.xlsx', data_only=True)

ws = wb['Acoes']
posicoes_acoes = []
for row in ws.iter_rows(min_row=2, values_only=True):
    produto, instituicao, conta, codigo, cnpj, isin, tipo, escriturador, qtd, qtd_disp, qtd_indisp, motivo, preco_fechamento, valor_atualizado = row
    if codigo and isinstance(codigo, str) and re.match(r'^[A-Z]{4}\d{1,2}$', codigo):
        posicoes_acoes.append({
            'ticker': codigo,
            'produto': produto,
            'quantidade': float(qtd) if qtd else 0,
            'preco_fechamento': float(preco_fechamento) if preco_fechamento else 0,
            'valor_atualizado': float(valor_atualizado) if valor_atualizado else 0
        })

# Fundo de Investimento sheet
posicoes_fiis = []
ws = wb['Fundo de Investimento']
for row in ws.iter_rows(min_row=2, values_only=True):
    produto, instituicao, conta, codigo, cnpj_fundo, isin, tipo, admin, qtd, qtd_disp, qtd_indisp, motivo, preco_fechamento, valor_atualizado = row
    if codigo and isinstance(codigo, str) and codigo.endswith('11'):
        posicoes_fiis.append({
            'ticker': codigo,
            'produto': produto,
            'quantidade': float(qtd) if qtd else 0,
            'preco_fechamento': float(preco_fechamento) if preco_fechamento else 0,
            'valor_atualizado': float(valor_atualizado) if valor_atualizado else 0
        })

# Renda Fixa sheet
posicoes_rf = []
ws = wb['Renda Fixa']
for row in ws.iter_rows(min_row=2, values_only=True):
    produto, instituicao, emissor, codigo, indexador, regime, emissao, vencimento, qtd, qtd_disp, qtd_indisp, motivo, contraparte, preco_mtm, valor_mtm, preco_curva, valor_curva, preco_fechamento, valor_fechamento = row
    if produto and 'CDB' in str(produto):
        posicoes_rf.append({
            'produto': produto,
            'codigo': codigo,
            'indexador': indexador,
            'vencimento': vencimento,
            'quantidade': float(qtd) if qtd else 0,
            'preco_mtm': float(preco_mtm) if preco_mtm else 0,
            'valor_mtm': float(valor_mtm) if valor_mtm else 0
        })

print(f'Ações: {len(posicoes_acoes)}, FIIs: {len(posicoes_fiis)}, RF: {len(posicoes_rf)}')

# Save all parsed data
data = {
    'renda_fixa': renda_fixa_positions,
    'renda_variavel': renda_var_positions,
    'movimentacoes': movimentacoes,
    'posicoes_acoes': posicoes_acoes,
    'posicoes_fiis': posicoes_fiis,
    'posicoes_rf': posicoes_rf,
    'extracted_at': datetime.now().isoformat()
}

with open('parsed_data.json', 'w', encoding='utf-8') as f:
    json.dump(data, f, ensure_ascii=False, indent=2)

print('Data saved to parsed_data.json')